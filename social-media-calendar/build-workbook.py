#!/usr/bin/env python3
"""Build a single styled .xlsx workbook from the 8 CSV tabs.
Upload the resulting file to Google Drive and open with Google Sheets."""
import csv
import glob
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))

# CSV file -> (Tab name, has_header_row)
TABS = [
    ("1-Dashboard.csv", "🏠 Dashboard", False),
    ("2-Content-Calendar.csv", "📅 Calendar", True),
    ("3-Content-Pillars-Ideas-Bank.csv", "💡 Ideas Bank", True),
    ("4-Raw-Assets-Library.csv", "🎞 Raw Assets", True),
    ("5-Festival-Calendar.csv", "🪔 Festivals", True),
    ("6-Roles-Workflow.csv", "👥 Roles & Workflow", False),
    ("7-Metrics-Tracker.csv", "📈 Metrics", True),
    ("8-Caption-Hashtag-Bank.csv", "# Caption & Hashtag Bank", True),
]

# Brand-ish palette
HEADER_FILL = PatternFill("solid", fgColor="1F2937")   # slate/charcoal
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
TITLE_FONT = Font(color="1F2937", bold=True, size=14)
ALT_FILL = PatternFill("solid", fgColor="F3F4F6")       # light grey zebra
THIN = Side(style="thin", color="D1D5DB")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

STATUS_COLORS = {
    "Idea": "E5E7EB",
    "Scripting": "FEF3C7",
    "Shooting": "FED7AA",
    "Editing": "BFDBFE",
    "Ready": "BBF7D0",
    "Scheduled": "DDD6FE",
    "Posted": "86EFAC",
    "Needs Rework": "FECACA",
}


def read_csv(path):
    with open(path, newline="", encoding="utf-8") as f:
        return [row for row in csv.reader(f)]


def build():
    wb = Workbook()
    wb.remove(wb.active)

    for fname, tabname, has_header in TABS:
        rows = read_csv(os.path.join(HERE, fname))
        ws = wb.create_sheet(title=tabname[:31])

        # find status column index if present
        status_col = None
        if has_header and rows:
            for i, h in enumerate(rows[0]):
                if h.strip().lower() == "status":
                    status_col = i
                    break

        max_cols = max((len(r) for r in rows), default=1)

        for r_idx, row in enumerate(rows, start=1):
            for c_idx in range(max_cols):
                val = row[c_idx] if c_idx < len(row) else ""
                cell = ws.cell(row=r_idx, column=c_idx + 1, value=val)
                cell.alignment = Alignment(vertical="top", wrap_text=True)

                if has_header and r_idx == 1:
                    cell.fill = HEADER_FILL
                    cell.font = HEADER_FONT
                    cell.alignment = Alignment(vertical="center", wrap_text=True)
                    cell.border = BORDER
                else:
                    # zebra striping for data rows
                    if has_header and r_idx % 2 == 0:
                        if not cell.fill or cell.fill.fgColor.rgb in (None, "00000000"):
                            cell.fill = ALT_FILL
                    # status color coding
                    if status_col is not None and c_idx == status_col and val.strip() in STATUS_COLORS:
                        cell.fill = PatternFill("solid", fgColor=STATUS_COLORS[val.strip()])
                    # bold the first-column labels on non-header sheets (Dashboard/Roles)
                    if not has_header and c_idx == 0 and val.strip():
                        cell.font = Font(bold=True, color="1F2937")

        # Title styling on non-header sheets first row
        if not has_header and rows:
            ws.cell(row=1, column=1).font = TITLE_FONT

        # Freeze header row / first data area
        ws.freeze_panes = "A2"

        # Column widths (sensible caps)
        for c_idx in range(max_cols):
            col_letter = get_column_letter(c_idx + 1)
            longest = 0
            for row in rows:
                if c_idx < len(row):
                    longest = max(longest, len(row[c_idx]))
            ws.column_dimensions[col_letter].width = min(max(12, longest + 2), 48)

        ws.sheet_view.showGridLines = False

    out = os.path.join(HERE, "Slate-Concepts-Social-Command-Center.xlsx")
    wb.save(out)
    print("Saved:", out)


if __name__ == "__main__":
    build()
